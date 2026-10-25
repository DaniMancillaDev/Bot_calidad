"""Regression checks for compatible capabilities recovered from PROD."""
import io
import hashlib
import tempfile
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings

from calidad.models import EvidenciaFotografica, ImportacionLote, RegistroDefecto
from calidad.services.importacion import (
    commit_importacion, validar_importacion_sin_ambiguedades,
)


class ImportRecoveryTests(TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        settings = override_settings(MEDIA_ROOT=self.root, FOTOS_ROOT=self.root / 'fotos')
        settings.enable()
        self.addCleanup(settings.disable)
        self.user = User.objects.create_user('import-recovery')
        self.client.force_login(self.user)

    def session(self, sid, contents=(b'original-one', b'original-two')):
        folder = self.root / 'importaciones_temp' / sid
        folder.mkdir(parents=True)
        for number, content in enumerate(contents, 1):
            (folder / f'{number}.jpg').write_bytes(content)
        return {
            'session_id': sid,
            'fotos': {str(n): {'filename': f'{n}.jpg', 'original_name': f'{n}.jpg'} for n in (1, 2)},
            'registros': [{'id': 'row', 'fotos_rango': '1-2', 'fotos_asignadas': ['2', '1']}],
        }

    def test_ambiguous_ranges_and_zip_numbers_are_rejected(self):
        with self.assertRaisesRegex(ValueError, 'filas'):
            validar_importacion_sin_ambiguedades([
                {'id': 'a', 'fotos_rango': '1-3'}, {'id': 'b', 'fotos_rango': '3-4'},
            ], {})
        with self.assertRaisesRegex(ValueError, 'ZIP'):
            validar_importacion_sin_ambiguedades([], {
                'a': {'original_name': '01.jpg'}, 'b': {'original_name': '1.png'},
            })
        validar_importacion_sin_ambiguedades([{'fotos_rango': '1, 1-2'}], {
            'a': {'original_name': '1.jpg'}, 'b': {'filename': '2.jpg'},
        })

    def test_upload_warns_before_creating_workspace(self):
        archive = io.BytesIO()
        with zipfile.ZipFile(archive, 'w') as output:
            output.writestr('one/01.jpg', b'first')
            output.writestr('two/01.jpg', b'second')
        with patch('calidad.views_importacion.parsear_txt', return_value=[]):
            response = self.client.post('/importacion/', {
                'data_file': SimpleUploadedFile('data.txt', b'test'),
                'zip_file': SimpleUploadedFile('photos.zip', archive.getvalue()),
            }, follow=True)
        self.assertContains(response, 'Importacion ambigua')
        self.assertNotIn('importacion_activa', self.client.session)
        self.assertEqual(RegistroDefecto.objects.count(), 0)

    def test_confirm_checks_ambiguity_and_missing_photos_before_writes(self):
        state = self.session('invalid')
        state['registros'].append(dict(state['registros'][0], id='second'))
        with self.assertRaisesRegex(ValueError, 'ambigua'):
            commit_importacion(state, self.user)
        state['registros'].pop()
        state['registros'][0]['fotos_asignadas'] = ['1']
        with self.assertRaisesRegex(ValueError, 'Faltan'):
            commit_importacion(state, self.user)
        self.assertEqual(ImportacionLote.objects.count(), 0)
        self.assertEqual(EvidenciaFotografica.objects.count(), 0)

    @patch('calidad.tasks.generar_thumbnails_lote_task.delay')
    def test_repeated_names_preserve_both_imports_and_selected_order(self, enqueue):
        first = commit_importacion(self.session('first'), self.user)
        before = {e.ruta_archivo: (self.root / e.ruta_archivo).read_bytes()
                  for e in EvidenciaFotografica.objects.all()}
        second = commit_importacion(self.session('second', (b'new-one', b'new-two')), self.user)
        self.assertNotEqual(first.id, second.id)
        self.assertEqual(EvidenciaFotografica.objects.count(), 4)
        self.assertEqual(len(set(EvidenciaFotografica.objects.values_list('ruta_archivo', flat=True))), 4)
        for name, content in before.items():
            self.assertEqual((self.root / name).read_bytes(), content)
        for lote in (first, second):
            evidence = list(EvidenciaFotografica.objects.filter(registro__importacion=lote).order_by('orden'))
            self.assertTrue(Path(evidence[0].ruta_archivo).name.startswith('2__imp_'))
            self.assertEqual([e.orden for e in evidence], [1, 2])
            self.assertEqual([e.es_portada for e in evidence], [True, False])
        self.assertEqual(enqueue.call_count, 2)

    @patch('calidad.tasks.generar_thumbnails_lote_task.delay')
    def test_even_generated_name_collision_cannot_overwrite(self, enqueue):
        with patch('calidad.services.importacion.uuid.uuid4', return_value=SimpleNamespace(hex='fixed')):
            commit_importacion(self.session('samepref-one'), self.user)
            state = self.session('samepref-two', (b'changed-one', b'changed-two'))
            before = {e.ruta_archivo: (self.root / e.ruta_archivo).read_bytes()
                      for e in EvidenciaFotografica.objects.all()}
            with self.assertRaises(FileExistsError):
                commit_importacion(state, self.user)
        self.assertEqual(ImportacionLote.objects.count(), 1)
        for path, content in before.items():
            self.assertEqual((self.root / path).read_bytes(), content)
        self.assertEqual(enqueue.call_count, 1)


class ExportRecoveryTests(TestCase):
    def test_export_keeps_all_photos_and_unique_names_across_groups_and_runs(self):
        from openpyxl import load_workbook
        from PIL import Image
        from pptx import Presentation
        from calidad.tasks import generar_excel_task

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            records = []
            for part in ('PN-1', 'PN-2'):
                record = RegistroDefecto.objects.create(
                    user_id=123, modelo='MODEL', numero_parte=part,
                    descripcion='DENT', responsable='XM', linea='T01', cantidad=2, fotos='',
                )
                records.append(record)
                for number in range(1, 21):
                    relative = f'fotos/123/{number:03d}_{record.id}.jpg'
                    path = root / relative
                    path.parent.mkdir(parents=True, exist_ok=True)
                    Image.new('RGB', (16, 24), 'red').save(path)
                    EvidenciaFotografica.objects.create(registro=record, ruta_archivo=relative, orden=number)
            before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in root.rglob('*.jpg')}
            translator = SimpleNamespace(translate=lambda *args: {'defect_en': 'DENT'})
            names_by_run = []
            with override_settings(MEDIA_ROOT=root), patch.object(generar_excel_task, 'update_state'), \
                    patch('shared.infrastructure.ai.defect_translator.build_translator', return_value=translator):
                for _ in range(2):
                    result = generar_excel_task.run([r.id for r in records], {}, str(root / 'fotos'), 'test')
                    with zipfile.ZipFile(result['file_path']) as archive:
                        pptx_names = [name for name in archive.namelist() if name.endswith('.pptx')]
                        self.assertEqual(len(pptx_names), 2)
                        self.assertTrue(any('PN1' in name.replace('-', '') for name in pptx_names))
                        photos = 0
                        for name in pptx_names:
                            presentation = Presentation(io.BytesIO(archive.read(name)))
                            photos += sum(len(slide.shapes) for slide in presentation.slides)
                        self.assertEqual(photos, 40)
                        workbook = load_workbook(io.BytesIO(archive.read('reporte_XM.xlsx')))
                        self.assertTrue(workbook.sheetnames)
                        names_by_run.append(set(pptx_names))
            self.assertTrue(names_by_run[0].isdisjoint(names_by_run[1]))
            self.assertEqual(before, {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in before})
            self.assertEqual(EvidenciaFotografica.objects.count(), 40)
