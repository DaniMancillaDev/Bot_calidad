import hashlib
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from zipfile import ZipFile

from django.conf import settings
from django.test import TestCase, override_settings
from openpyxl import load_workbook
from PIL import Image

from calidad.models import RegistroDefecto, EvidenciaFotografica
from calidad.tasks import generar_excel_box_condition_task


class BoxTemplateRecoveryTests(TestCase):
    def test_real_task_uses_shipped_template_keeps_diagram_and_all_evidence(self):
        template = Path(settings.BASE_DIR).parent / 'Daily Tracking Box Condition.xlsx'
        before_template = hashlib.sha256(template.read_bytes()).hexdigest()
        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            records = []
            for line, quantity in (('T01', 10), ('T03', 15)):
                record = RegistroDefecto.objects.create(
                    user_id=123, modelo='MODEL', numero_parte='PN', descripcion='CARTON DAMAGED',
                    responsable='XM', linea=line, cantidad=quantity, fotos='',
                )
                records.append(record)
                for number in range(9):
                    path = root / f'{record.id}_{number}.jpg'
                    Image.new('RGB', (24, 16), 'blue').save(path)
                    EvidenciaFotografica.objects.create(registro=record, ruta_archivo=path.name, orden=number+1)
            before = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in root.glob('*.jpg')}
            translator = SimpleNamespace(translate=lambda *args: {'defect_en': 'CARTON DAMAGED'})
            with override_settings(MEDIA_ROOT=root), \
                    patch.object(generar_excel_box_condition_task, 'update_state'), \
                    patch('shared.infrastructure.ai.defect_translator.build_translator', return_value=translator):
                result = generar_excel_box_condition_task.run(None, str(root / 'fotos'), registros_ids=[r.id for r in records])
            self.assertTrue(result['filename'].endswith('.xlsx'))
            output = load_workbook(result['file_path'])
            self.assertEqual(output.sheetnames[:4], load_workbook(template).sheetnames)
            report = output['23-09-26']
            self.assertEqual(report.cell(2, 8).value, 'T01/T03')
            self.assertEqual(report.cell(2, 9).value, 25)
            self.assertEqual(len(output['Fotos Box Condition']._images), 18)
            self.assertEqual({output['Fotos Box Condition'].cell(row, 2).value for row in range(2, 20)},
                             {record.id for record in records})
            self.assertEqual(str(report.print_area), "'23-09-26'!$A$1:$P$2")
            with ZipFile(template) as original, ZipFile(result['file_path']) as generated:
                for name in ('xl/drawings/templateDrawing2.xml', 'xl/drawings/templateDrawing3.xml',
                             'xl/worksheets/sheet1.xml', 'xl/worksheets/sheet2.xml', 'xl/worksheets/sheet3.xml'):
                    self.assertEqual(original.read(name), generated.read(name), name)
            self.assertEqual(before, {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in before})
            self.assertEqual(before_template, hashlib.sha256(template.read_bytes()).hexdigest())
