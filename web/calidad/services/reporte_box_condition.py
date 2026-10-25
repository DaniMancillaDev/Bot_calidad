import os
import io
import shutil
import tempfile
import zipfile
from xml.etree import ElementTree as ET
from pathlib import Path
from PIL import Image, ImageOps
import openpyxl
from openpyxl.drawing.image import Image as OpenpyxlImage
from openpyxl.styles import Alignment, Font
from openpyxl.drawing.xdr import XDRPoint2D, XDRPositiveSize2D
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker

DUTY_BOX_CONDITION = {
    "XM": "XM",
    "WH": "MAT WH",
    "TSCEM": "PROD",
    "MNT": "MNT",
    "MNTO": "MNTO",
}

TURNO_SHIFT = {
    "A": "First",
    "B": "Second",
    "C": "Third",
}

def es_box_condition(defecto, translator):
    """
    Filtro aislado para defectos Box Condition.
    Pasa por el traductor canónico y verifica palabras clave.
    """
    if not defecto:
        return False
    # Para ser robusto, usamos el mismo traductor del reporte principal
    # pero sin alterarlo.
    from shared.infrastructure.ai.defect_translator import build_translator
    if translator is None:
        translator = build_translator()
        
    translation = translator.translate(defecto, "")
    canonico = translation['defect_en'].upper() if translation else str(defecto).upper()
    keywords = ["CAJA", "CARTON", "BOX"]
    for kw in keywords:
        if kw in canonico or kw in str(defecto).upper():
            return True
    return False

def pixels_to_EMU(pixels):
    return int(pixels * 914400 / 96)

import logging
logger = logging.getLogger(__name__)


def _preserve_auxiliary_template_parts(template_path, output_path, edited_sheet):
    """openpyxl cannot retain the process diagram's native Excel shapes."""
    ns = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
    rel_ns = 'http://schemas.openxmlformats.org/package/2006/relationships'
    with zipfile.ZipFile(template_path) as archive:
        original = {name: archive.read(name) for name in archive.namelist()}
    drawings = {name: content for name, content in original.items()
                if name.startswith('xl/drawings/templateDrawing')}
    if not drawings:
        return
    restored = dict(drawings)
    sheets = ET.fromstring(original['xl/workbook.xml']).find(f'{{{ns}}}sheets')
    for index, sheet in enumerate(sheets, 1):
        if sheet.get('name') in (edited_sheet, 'Fotos Box Condition'):
            continue
        name = f'xl/worksheets/sheet{index}.xml'
        restored[name] = original[name]
        relation = f'xl/worksheets/_rels/sheet{index}.xml.rels'
        if relation in original:
            restored[relation] = original[relation]
    restored['xl/sharedStrings.xml'] = original['xl/sharedStrings.xml']
    with zipfile.ZipFile(output_path) as archive:
        output = {name: archive.read(name) for name in archive.namelist()}
    output.update(restored)
    relationships = ET.fromstring(output['xl/_rels/workbook.xml.rels'])
    ET.SubElement(relationships, f'{{{rel_ns}}}Relationship', {
        'Id': 'rIdBoxTemplateStrings', 'Type': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings',
        'Target': 'sharedStrings.xml',
    })
    output['xl/_rels/workbook.xml.rels'] = ET.tostring(relationships, encoding='utf-8', xml_declaration=True)
    content_types = ET.fromstring(output['[Content_Types].xml'])
    existing = {entry.get('PartName') for entry in content_types}
    for entry in ET.fromstring(original['[Content_Types].xml']):
        if entry.get('PartName', '').lstrip('/') in restored and entry.get('PartName') not in existing:
            content_types.append(entry)
    output['[Content_Types].xml'] = ET.tostring(content_types, encoding='utf-8', xml_declaration=True)
    # Only the newly generated report is replaced; the template remains read-only.
    with tempfile.NamedTemporaryFile(dir=Path(output_path).parent, suffix='.xlsx') as temporary:
        with zipfile.ZipFile(temporary, 'w', zipfile.ZIP_DEFLATED) as archive:
            for name, content in output.items():
                archive.writestr(name, content)
        temporary.flush()
        shutil.copyfile(temporary.name, output_path)

def box_condition_group_key(r, duty, translator):
    # ponytail: La clave debe incluir el responsable para no cruzar Cajas Dañadas de XM y WH
    # No usamos fecha en la clave para consolidar ocurrencias de distintos días en el mismo grupo
    from shared.infrastructure.ai.defect_translator import build_translator
    if not translator:
        translator = build_translator()
    
    modelo = (r.modelo or "").strip().upper()
    pn = (r.numero_parte or "").strip().upper()
    
    # Normalización del defecto usando el catálogo existente
    translation = translator.translate(r.descripcion, "")
    defecto = translation['defect_en'].upper() if translation else str(r.descripcion).strip().upper()
    
    return (modelo, pn, defecto, duty)

def generate_box_condition_excel(registros_orm, template_path, fotos_dir, output_path, translator=None):
    """
    Genera el reporte aislado Daily Tracking Box Condition.
    Una fila por modelo, PN, defecto y DUTY, con las lineas combinadas.
    """
    from datetime import datetime
    
    # 1. Copiar template original
    shutil.copy(template_path, output_path)
    
    wb = openpyxl.load_workbook(output_path)
    # Buscar la hoja correcta (la que tiene FIRTS/THIRD)
    target_ws = None
    for sn in wb.sheetnames:
        ws = wb[sn]
        first_row = [c.value for c in ws[1]]
        if first_row and any("FIRTS/THIRD" in str(v) for v in first_row):
            target_ws = ws
            break
            
    if not target_ws:
        # Fallback a nombre de hoja si alguien la cambió
        target_ws = wb['23-09-26'] if '23-09-26' in wb.sheetnames else wb.active
        
    ws = target_ws
    fotos_ws = wb['Fotos Box Condition'] if 'Fotos Box Condition' in wb.sheetnames else wb.create_sheet('Fotos Box Condition')
    fotos_ws._images.clear()
    fotos_ws.delete_rows(1, fotos_ws.max_row)
    fotos_ws.append(['Grupo', 'Registro ID', 'Modelo', 'PN', 'Linea', 'Defecto', 'Duty', 'Archivo', 'Evidencia'])
    fotos_ws.column_dimensions['A'].width = 12
    fotos_ws.column_dimensions['B'].width = 12
    fotos_ws.column_dimensions['C'].width = 18
    fotos_ws.column_dimensions['D'].width = 18
    fotos_ws.column_dimensions['E'].width = 14
    fotos_ws.column_dimensions['F'].width = 28
    fotos_ws.column_dimensions['G'].width = 14
    fotos_ws.column_dimensions['H'].width = 28
    fotos_ws.column_dimensions['I'].width = 28
    image_buffers = []



    # Determinar fecha global del reporte (puede venir del contexto del primer registro o la fecha actual)
    fecha_reporte = None
    if registros_orm:
        primer_lote = next((r.importacion for r in registros_orm if r.importacion), None)
        if primer_lote and primer_lote.fecha_inicio:
            fecha_reporte = primer_lote.fecha_inicio.date()
        elif registros_orm[0].fecha_registro:
            fecha_reporte = registros_orm[0].fecha_registro.date()
    if not fecha_reporte:
        fecha_reporte = datetime.now().date()

    # 2. Agrupar registros
    grupos = {}
    for r in registros_orm:
        if not es_box_condition(r.descripcion, translator):
            continue
            
        responsable_db = (r.responsable or "").strip().upper()
        if responsable_db not in DUTY_BOX_CONDITION:
            continue
            
        duty_mapped = DUTY_BOX_CONDITION[responsable_db]
        
        clave = box_condition_group_key(r, duty_mapped, translator)
        if clave not in grupos:
            grupos[clave] = {
                'modelo': clave[0],
                'pn': clave[1],
                'linea': '',
                'lineas': [],
                'defecto': clave[2],
                'duty': clave[3],
                'cantidad': 0,
                'turno': (r.turno or "").strip().upper(),
                'registros': []
            }
            
        linea = (r.linea or '').strip().upper()
        if linea and linea not in grupos[clave]['lineas']:
            grupos[clave]['lineas'].append(linea)
        grupos[clave]['linea'] = '/'.join(grupos[clave]['lineas'])
        grupos[clave]['cantidad'] += int(r.cantidad) if r.cantidad else 1
        grupos[clave]['registros'].append(r)

    # 3. Renderizar Main Sheet y Photos Sheet
    ws._images.clear()

    # Los headers están en la fila 1
    START_ROW = 2
    current_row = START_ROW
    current_foto_row = 2

    CELL_WIDTH_PX = 100
    CELL_HEIGHT_PX = 100
    MARGIN_PX = 5
    
    FOTO_WIDTH_PX = 300
    FOTO_HEIGHT_PX = 300

    for group_idx, (clave, grupo) in enumerate(grupos.items(), start=1):
        group_id_str = f"GRP-{group_idx}"
        
        # --- MAIN SHEET ---
        ws.cell(row=current_row, column=1).value = fecha_reporte.strftime("%Y-%m-%d")
        ws.cell(row=current_row, column=2).value = grupo['modelo']
        shift_val = TURNO_SHIFT.get(grupo['turno'], "")
        ws.cell(row=current_row, column=3).value = shift_val
        ws.cell(row=current_row, column=4).value = grupo['pn']
        ws.cell(row=current_row, column=8).value = grupo['linea']
        ws.cell(row=current_row, column=9).value = grupo['cantidad']
        ws.cell(row=current_row, column=10).value = grupo['defecto']
        ws.cell(row=current_row, column=13).value = grupo['duty']
        
        # ponytail: Aplicar formato original (centrado, wrap_text y negrita) a las celdas de la fila
        align = Alignment(horizontal='center', vertical='center', wrap_text=True)
        font = Font(bold=True)
        for col_i in range(1, 15):
            c = ws.cell(row=current_row, column=col_i)
            c.alignment = align
            c.font = font
        
        evidencias_grupo = []
        for r in grupo['registros']:
            for ev in r.evidencias_v2.all():
                if not (isinstance(ev.metadatos, dict) and ev.metadatos.get('excluida')):
                    evidencias_grupo.append(ev)

        evidencias_portada = evidencias_grupo[:15]
                
        if evidencias_portada:
            from django.conf import settings
            x_offset = MARGIN_PX
            max_h = 0
            
            for ev in evidencias_portada:
                ev_path_str = ev.ruta_archivo
                if ev_path_str.startswith('media_files/'): ev_path_str = ev_path_str[12:]
                elif ev_path_str.startswith('/media_files/'): ev_path_str = ev_path_str[13:]
                foto_path = Path(settings.MEDIA_ROOT) / ev_path_str.lstrip('/')
                
                if foto_path.exists():
                    try:
                        with Image.open(foto_path) as im:
                            im = ImageOps.exif_transpose(im)
                            orig_w, orig_h = im.size
                            # Tamaño mucho más grande (200x200) para claridad y zoom
                            TARGET_SIZE = 200
                            ratio = min(TARGET_SIZE / orig_w, TARGET_SIZE / orig_h)
                            disp_w = int(orig_w * ratio)
                            disp_h = int(orig_h * ratio)
                            
                            # SCALE_FACTOR = 2 (doble resolución física para nitidez extrema)
                            phys_w = disp_w * 2
                            phys_h = disp_h * 2
                            
                            im_resized = im.resize((phys_w, phys_h), Image.Resampling.LANCZOS)
                            if ev.angulo_rotacion:
                                im_resized = im_resized.rotate(-ev.angulo_rotacion, expand=True)
                            
                            im_resized = im_resized.convert('RGB')
                            img_byte_arr = io.BytesIO()
                            image_buffers.append(img_byte_arr)
                            im_resized.save(img_byte_arr, format='JPEG', quality=85, optimize=True)
                            img_byte_arr.seek(0)
                            
                            xl_img = OpenpyxlImage(img_byte_arr)
                            col_idx_0based = 11  # Col L (12) es defective picture
                            row_idx_0based = current_row - 1
                            y_offset = MARGIN_PX
                            
                            marker = AnchorMarker(col=col_idx_0based, colOff=pixels_to_EMU(x_offset), row=row_idx_0based, rowOff=pixels_to_EMU(y_offset))
                            # El ancla usa el tamaño visual, pero la imagen insertada tiene el doble de píxeles
                            size = XDRPositiveSize2D(pixels_to_EMU(disp_w), pixels_to_EMU(disp_h))
                            xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
                            ws.add_image(xl_img)
                            
                            x_offset += disp_w + MARGIN_PX
                            max_h = max(max_h, disp_h)
                    except Exception as e:
                        logger.exception(f"Error procesando imagen {foto_path} para la hoja principal: {e}")
            
            if max_h > 0:
                # Altura en puntos (pixels * 0.75) para que quepa la imagen grande
                row_h_pts = int((max_h + 2 * MARGIN_PX) * 0.75)
                ws.row_dimensions[current_row].height = max(80, row_h_pts)
                
                col_width = max(15, int(x_offset / 7))
                # Ajustar el ancho de la columna si la tira es más larga (solo lo ensanchará para el más largo)
                if ws.column_dimensions['L'].width is None or ws.column_dimensions['L'].width < col_width:
                    ws.column_dimensions['L'].width = col_width

        for ev in evidencias_grupo:
            from django.conf import settings
            ev_path_str = ev.ruta_archivo
            if ev_path_str.startswith('media_files/'): ev_path_str = ev_path_str[12:]
            elif ev_path_str.startswith('/media_files/'): ev_path_str = ev_path_str[13:]
            foto_path = Path(settings.MEDIA_ROOT) / ev_path_str.lstrip('/')
            if not foto_path.exists():
                continue

            fotos_ws.cell(row=current_foto_row, column=1).value = group_id_str
            fotos_ws.cell(row=current_foto_row, column=2).value = ev.registro_id
            fotos_ws.cell(row=current_foto_row, column=3).value = grupo['modelo']
            fotos_ws.cell(row=current_foto_row, column=4).value = grupo['pn']
            fotos_ws.cell(row=current_foto_row, column=5).value = grupo['linea']
            fotos_ws.cell(row=current_foto_row, column=6).value = grupo['defecto']
            fotos_ws.cell(row=current_foto_row, column=7).value = grupo['duty']
            fotos_ws.cell(row=current_foto_row, column=8).value = foto_path.name
            fotos_ws.row_dimensions[current_foto_row].height = 120

            try:
                with Image.open(foto_path) as im:
                    im = ImageOps.exif_transpose(im)
                    if ev.angulo_rotacion:
                        im = im.rotate(-ev.angulo_rotacion, expand=True)
                    orig_w, orig_h = im.size
                    ratio = min(150 / orig_w, 150 / orig_h)
                    disp_w = max(1, int(orig_w * ratio))
                    disp_h = max(1, int(orig_h * ratio))
                    im_resized = im.resize((disp_w, disp_h), Image.Resampling.LANCZOS).convert('RGB')
                    img_byte_arr = io.BytesIO()
                    image_buffers.append(img_byte_arr)
                    im_resized.save(img_byte_arr, format='JPEG', quality=85, optimize=True)
                    img_byte_arr.seek(0)

                    xl_img = OpenpyxlImage(img_byte_arr)
                    marker = AnchorMarker(col=8, colOff=pixels_to_EMU(MARGIN_PX), row=current_foto_row - 1, rowOff=pixels_to_EMU(MARGIN_PX))
                    size = XDRPositiveSize2D(pixels_to_EMU(disp_w), pixels_to_EMU(disp_h))
                    xl_img.anchor = OneCellAnchor(_from=marker, ext=size)
                    fotos_ws.add_image(xl_img)
            except Exception as e:
                logger.exception(f"Error procesando imagen {foto_path} para hoja de evidencias: {e}")

            current_foto_row += 1
        
        current_row += 1
        
    ws.print_area = f'A1:P{max(2, current_row - 1)}'
    wb.save(output_path)
    _preserve_auxiliary_template_parts(template_path, output_path, ws.title)
