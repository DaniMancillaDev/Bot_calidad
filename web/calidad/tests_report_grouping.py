import tempfile
from pathlib import Path
from types import SimpleNamespace

from django.test import SimpleTestCase, TestCase, override_settings
from openpyxl import load_workbook
from PIL import Image

from calidad.models import RegistroDefecto, EvidenciaFotografica
from calidad.services.report_grouping import DEV_DESCRIPTION, PROD_MODEL_RESPONSIBLE, report_group_key
from calidad.tasks import _consolidar_registros


class GroupingComparisonTests(SimpleTestCase):
    def test_consolidated_group_keeps_distinct_invoices(self):
        rows = [
            {'id': 1, 'numero_parte': 'PN', 'descripcion': 'DENT', 'invoice': 'INV-1'},
            {'id': 2, 'numero_parte': 'PN', 'descripcion': 'DENT', 'invoice': 'INV-2'},
            {'id': 3, 'numero_parte': 'PN', 'descripcion': 'DENT', 'invoice': 'INV-1'},
        ]

        groups = _consolidar_registros(rows)

        self.assertEqual(groups[0]['invoice'], 'INV-1\nINV-2')

    def test_both_rules_keep_quantities_record_ids_and_photos(self):
        base = {'id': 1, 'numero_parte': 'PN', 'modelo': 'MODEL-A', 'responsable': 'WH',
                'descripcion': 'DENT', 'cantidad': 2, 'evidencias': ['photo-a']}
        cases = [('numero_parte', 'OTHER', 2, 2), ('modelo', 'MODEL-B', 1, 2),
                 ('responsable', 'TSCEM', 1, 2), ('descripcion', 'SCRATCH', 2, 1),
                 ('cantidad', 3, 1, 1)]
        for field, value, dev_count, prod_count in cases:
            second = dict(base, id=2, cantidad=3, evidencias=['photo-b'])
            second[field] = value
            for rule, count in ((DEV_DESCRIPTION, dev_count), (PROD_MODEL_RESPONSIBLE, prod_count)):
                with self.subTest(field=field, rule=rule):
                    groups = _consolidar_registros([base, second], grouping_rule=rule)
                    self.assertEqual(len(groups), count)
                    self.assertEqual(sum(g['cantidad'] for g in groups), 5)
                    self.assertEqual(sorted(i for g in groups for i in g['ids_originales']), [1, 2])
                    self.assertEqual(sorted(p for g in groups for p in g['evidencias']), ['photo-a', 'photo-b'])
                    for group in groups:
                        self.assertEqual(set(group['descripciones_originales']),
                                         {r['descripcion'] for r in group['registros']})
        self.assertEqual(_consolidar_registros([base])[0]['key'], ('PN', 'DENT'))

    def test_rule_requires_an_explicit_supported_name(self):
        with self.assertRaises(ValueError):
            report_group_key({}, 'unreviewed')
        self.assertEqual(report_group_key({'numero_parte': ' pn ', 'descripcion': ' dent '}), ('PN', 'DENT'))


class GroupingExcelTests(TestCase):
    def test_each_rule_exports_all_its_photos_and_rejects_an_outside_record(self):
        from calidad.services.reporte_excel import generate_excel, PLANTILLA_PATH

        with tempfile.TemporaryDirectory() as folder:
            root = Path(folder)
            rows = []
            for description in ('DENT', 'SCRATCH', 'OUTSIDE'):
                r = RegistroDefecto.objects.create(user_id=123, modelo='MODEL', numero_parte='PN',
                    responsable='XM', descripcion=description, linea='T01', cantidad=2, fotos='')
                p = root / f'{r.id}.jpg'
                Image.new('RGB', (24, 16), 'red').save(p)
                ev = EvidenciaFotografica.objects.create(registro=r, ruta_archivo=p.name, orden=1)
                rows.append({'id': r.id, 'modelo': r.modelo, 'numero_parte': r.numero_parte,
                    'responsable': r.responsable, 'descripcion': r.descripcion,
                    'cantidad': r.cantidad, 'evidencias': [ev]})
            translator = SimpleNamespace(translate=lambda *args: {'defect_en': 'DENT'})
            template_images = sum(len(s._images) for s in load_workbook(PLANTILLA_PATH)) if PLANTILLA_PATH.exists() else 0
            for rule in (DEV_DESCRIPTION, PROD_MODEL_RESPONSIBLE):
                groups = _consolidar_registros(rows[:2], translator, grouping_rule=rule)
                groups[0]['evidencias'].append(rows[2]['evidencias'][0])
                output = root / f'{rule}.xlsx'
                with override_settings(MEDIA_ROOT=root), self.assertLogs('calidad.services.reporte_excel', level='ERROR'):
                    generate_excel(groups, {}, root, output, translator=translator)
                workbook = load_workbook(output)
                self.assertEqual(sum(len(sheet._images) for sheet in workbook), template_images + 2)
                self.assertEqual(EvidenciaFotografica.objects.count(), 3)
