import hashlib
import importlib.util
import io
import posixpath
import unittest
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

ROOT = Path(__file__).parents[2]
TEMPLATE = ROOT / 'Daily Tracking Box Condition.xlsx'
spec = importlib.util.spec_from_file_location('sanitize', ROOT / 'scripts/sanitize_box_template.py')
sanitize = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sanitize)


class BoxTemplateChecks(unittest.TestCase):
    def test_unknown_original_cannot_be_sanitized_by_guessing(self):
        with self.assertRaisesRegex(ValueError, 'Unreviewed'):
            sanitize.sanitized_bytes(b'unreviewed')

    def test_template_retains_structure_without_sample_data_or_photos(self):
        workbook = load_workbook(TEMPLATE)
        self.assertEqual(workbook.sheetnames, [
            'daily tracking box condition ', 'Gantt Chart', 'Process chart', '23-09-26',
        ])
        for name, first_row in (('daily tracking box condition ', 3), ('Gantt Chart', 2), ('23-09-26', 2)):
            for row in workbook[name].iter_rows(min_row=first_row):
                self.assertTrue(all(cell.value is None for cell in row))
        self.assertEqual(str(next(iter(workbook.worksheets[0].merged_cells.ranges))), 'K1:M1')
        self.assertEqual(workbook['23-09-26']['C1'].value, 'FIRTS/THIRD')
        self.assertEqual(sum(len(sheet._images) for sheet in workbook), 0)
        self.assertEqual(sum(cell.data_type == 'f' for sheet in workbook for row in sheet for cell in row), 0)
        self.assertEqual(sum(len(sheet.data_validations.dataValidation) for sheet in workbook), 0)
        self.assertEqual(len(workbook.defined_names), 0)
        with ZipFile(TEMPLATE) as archive:
            self.assertFalse(any(name.startswith(('xl/media/', 'customXml/')) or 'comments' in name
                                 for name in archive.namelist()))
            self.assertIn('xl/drawings/templateDrawing2.xml', archive.namelist())
            for name in archive.namelist():
                if not name.endswith('.rels'):
                    continue
                base = '' if name == '_rels/.rels' else name.rsplit('/_rels/', 1)[0]
                relations = ET.fromstring(archive.read(name))
                self.assertEqual(len({item.get('Id') for item in relations}), len(relations))
                for item in relations:
                    self.assertNotEqual(item.get('TargetMode'), 'External')
                    target = item.get('Target')
                    resolved = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(base, target))
                    self.assertIn(resolved, archive.namelist(), (name, target))
