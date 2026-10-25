"""Run an isolated image check without starting any application or attaching real data."""
import argparse
import json
import subprocess
import uuid
from pathlib import Path

from preflight_prod_dev import docker_command, image_sources

PROBE = '''
import hashlib,json
from pathlib import Path
from openpyxl import load_workbook
root=Path('/app')
files={}
for directory in ('web','shared','bot'):
 for p in (root/directory).rglob('*'):
  if p.is_file() and '__pycache__' not in p.parts and p.suffix in ('.py','.html','.css','.js'):
   files[p.relative_to(root).as_posix()]=hashlib.sha256(p.read_bytes()).hexdigest()
for name in ('pyproject.toml','uv.lock','plantilla_reporte.xlsx','Daily Tracking Box Condition.xlsx','shared/infrastructure/ai/defect_catalog.json'):
 files[name]=hashlib.sha256((root/name).read_bytes()).hexdigest()
assert not (root/'.env').exists() and not (root/'.env.local').exists()
workbook=load_workbook(root/'Daily Tracking Box Condition.xlsx')
assert len(workbook.sheetnames)==4 and workbook['23-09-26']['C1'].value=='FIRTS/THIRD'
assert sum(len(sheet._images) for sheet in workbook)==0
print(json.dumps(files,sort_keys=True))
'''


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--repo', type=Path, default=Path.cwd())
    parser.add_argument('--image', required=True)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    if args.output.exists():
        parser.error('Do not overwrite existing image evidence')
    docker = docker_command()
    image_id = subprocess.check_output(docker + ['image', 'inspect', '--format', '{{.Id}}', args.image], text=True).strip()
    name = 'iqa-image-check-' + uuid.uuid4().hex[:12]
    subprocess.run(docker + ['create', '--name', name, '--network', 'none',
        '--tmpfs', '/app/data', '--tmpfs', '/app/media_files', '--entrypoint', '/app/.venv/bin/python',
        image_id, '-c', PROBE], check=True, stdout=subprocess.DEVNULL)
    output = subprocess.check_output(docker + ['start', '-a', name], text=True)
    status = subprocess.check_output(docker + ['inspect', '--format', '{{.State.ExitCode}}', name], text=True).strip()
    if status != '0':
        raise RuntimeError(f'Image verification failed; inspect isolated container {name}')
    files = json.loads(output)
    if files != image_sources(args.repo):
        raise RuntimeError('Built image differs from current DEV sources/resources; rebuild it')
    result = {'image_id': image_id, 'files': files, 'verification_container': name}
    with args.output.open('x', encoding='utf-8') as target:
        json.dump(result, target, indent=2)
        target.write('\n')
    print(f'Verified {image_id}, {len(files)} source/resource hashes; container {name} is stopped')


if __name__ == '__main__':
    main()
